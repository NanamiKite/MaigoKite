"""Bound local document extraction, shared by attachments, tools and previews."""
from __future__ import annotations

import hashlib
import re
from collections.abc import Sequence
from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
from typing import Literal
from zipfile import BadZipFile, ZipFile

from anyio import BrokenWorkerProcess, CapacityLimiter, fail_after, to_process
from docx import Document
from docx.oxml.exceptions import InvalidXmlError as DocxInvalidXmlError
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from pptx import Presentation
from pptx.exc import InvalidXmlError as PptxInvalidXmlError
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pydantic import BaseModel
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from pypdf.generic import Destination

from agent.errors import UserFacingError

type DocumentFormat = Literal["pdf", "docx", "xlsx", "pptx"]
type DocumentKind = Literal["page", "paragraph", "table", "sheet", "slide"]
type OutlineEntry = Destination | Sequence[OutlineEntry]
DOCUMENT_LIMIT = 16 * 1024 * 1024
PARSER_VERSION = "text-v3"
_parser_slots = CapacityLimiter(2)


class DocumentUnit(BaseModel):
    kind: DocumentKind
    text: str


class DocumentLocation(BaseModel):
    unit: int
    kind: DocumentKind
    text: str
    text_offset: int
    level: int | None = None


class DocumentNavigation(BaseModel):
    format: DocumentFormat
    sha256: str
    parser_version: str
    action: Literal["outline", "search"]
    query: str
    total_units: int
    empty_text_units: int
    items: list[DocumentLocation]
    next_offset: int | None
    outline_truncated: bool


@dataclass(frozen=True)
class ParsedDocument:
    format: DocumentFormat
    units: tuple[DocumentUnit, ...]
    outline: tuple[DocumentLocation, ...]
    outline_truncated: bool = False


class DocumentPage(BaseModel):
    mode: Literal["text_layer", "ocr"] = "text_layer"
    format: DocumentFormat
    sha256: str
    parser_version: str
    unit: int
    end_unit: int
    total_units: int
    kind: DocumentKind
    text: str
    text_offset: int
    next_offset: int | None
    next_unit: int | None


def document_format(data: bytes) -> DocumentFormat | None:
    if data.startswith(b"%PDF-"):
        return "pdf"
    if not data.startswith(b"PK\x03\x04"):
        return None
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names:
                return None
            formats: tuple[tuple[str, DocumentFormat], ...] = (
                ("word/document.xml", "docx"), ("xl/workbook.xml", "xlsx"), ("ppt/presentation.xml", "pptx"),
            )
            for member, format in formats:
                if member in names:
                    return format
    except BadZipFile as exc:
        raise UserFacingError("文档已损坏，无法读取。", code="DOCUMENT_INVALID") from exc
    return None


@lru_cache(maxsize=4)
def _parse(data: bytes, format: DocumentFormat, parser_version: str) -> ParsedDocument:
    # This cache lives only in the worker: exact bytes + parser version, at most four bounded documents.
    units: list[DocumentUnit] = []
    outline: list[DocumentLocation] = []
    if format == "pdf":
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise UserFacingError("请先解密文档后再添加。", code="DOCUMENT_ENCRYPTED")
        if len(reader.pages) > 200:
            raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
        for page in reader.pages:
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > 4 * 1024 * 1024:
                raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
            units.append(DocumentUnit(kind="page", text=page.extract_text()))
            if sum(len(unit.text) for unit in units) > 1_000_000:
                raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
        pending: list[tuple[int, OutlineEntry]] = [(0, item) for item in reversed(reader.outline)]
        while pending and len(outline) < 1000:
            level, item = pending.pop()
            if not isinstance(item, Destination):
                pending.extend((level + 1, child) for child in reversed(item))
                continue
            page_number = reader.get_destination_page_number(item)
            if page_number is None or not 0 <= page_number < len(units):
                continue
            outline.append(DocumentLocation(
                unit=page_number + 1, kind="page", text=(item.title or "")[:200], text_offset=0, level=level,
            ))
        return ParsedDocument(format, tuple(units), tuple(outline), outline_truncated=bool(pending))
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if (len(entries) > 3000 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024
                or any(entry.file_size > 8 * 1024 * 1024 for entry in entries if entry.filename.endswith(".xml"))):
            raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
        if any(entry.flag_bits & 1 for entry in entries):
            raise UserFacingError("请先解密文档后再添加。", code="DOCUMENT_ENCRYPTED")
    if format == "xlsx":
        try:
            workbook = load_workbook(BytesIO(data), data_only=False, keep_links=False)
        except (KeyError, ValueError, OSError) as exc:
            # With an in-memory ZIP, these loader errors identify missing/invalid package parts.
            raise UserFacingError("文档已损坏，无法读取。", code="DOCUMENT_INVALID") from exc
        try:
            if len(workbook.worksheets) > 100:
                raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
            text_size = 0
            for sheet in workbook.worksheets:
                if sheet.max_row > 10_000 or sheet.max_column > 200:
                    raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
                outline.append(DocumentLocation(unit=len(units) + 1, kind="sheet", text=sheet.title,
                                                text_offset=0, level=0))
                for start in range(1, sheet.max_row + 1, 100):
                    rows = [f"Sheet: {sheet.title}; rows {start}-{min(start + 99, sheet.max_row)}"]
                    text_size += len(rows[0]) + 1
                    for row in sheet.iter_rows(min_row=start, max_row=min(start + 99, sheet.max_row)):
                        values: list[str] = []
                        for cell in row:
                            value = cell.value
                            if value is None:
                                continue
                            if isinstance(value, ArrayFormula):
                                value = str(value.text or "")
                            elif isinstance(value, DataTableFormula):
                                value = f"[data-table formula: {dict(value)}]"
                            values.append(f"{cell.coordinate}={value}")
                        rows.append("\t".join(values))
                        text_size += len(rows[-1]) + 1
                        if text_size > 1_000_000:
                            raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
                    units.append(DocumentUnit(kind="sheet", text="\n".join(rows)))
            return ParsedDocument(format, tuple(units), tuple(outline))
        finally:
            workbook.close()
    if format == "pptx":
        try:
            presentation = Presentation(BytesIO(data))
            slides = presentation.slides
        except (KeyError, ValueError) as exc:
            # Resolving slides also checks their relationship IDs, before text extraction starts.
            raise UserFacingError("文档已损坏，无法读取。", code="DOCUMENT_INVALID") from exc
        if len(slides) > 200:
            raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
        text_size = 0
        for number, slide in enumerate(slides, start=1):
            texts: list[str] = []
            shapes = list(reversed(slide.shapes))
            while shapes:
                shape = shapes.pop()
                if isinstance(shape, GroupShape):
                    shapes.extend(reversed(shape.shapes))
                    continue
                if isinstance(shape, Shape) and shape.has_text_frame:
                    texts.append(shape.text)
                elif isinstance(shape, GraphicFrame) and shape.has_table:
                    texts.append("\n".join("\t".join(cell.text for cell in row.cells) for row in shape.table.rows))
            text = "\n\n".join(texts)
            text_size += len(text)
            if text_size > 1_000_000:
                raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
            units.append(DocumentUnit(kind="slide", text=text))
            title = slide.shapes.title
            outline.append(DocumentLocation(unit=number, kind="slide", text=title.text[:200] if title else "",
                                            text_offset=0, level=0))
        return ParsedDocument(format, tuple(units), tuple(outline))
    try:
        document = Document(BytesIO(data))
    except (KeyError, ValueError) as exc:
        # The loader requires one main-document relationship and valid ZIP part/content-type mappings.
        raise UserFacingError("文档已损坏，无法读取。", code="DOCUMENT_INVALID") from exc
    text_size = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            units.append(DocumentUnit(kind="paragraph", text=block.text))
            # OOXML built-in headings use English style names even in localized Word installations.
            style = block.style
            style_name = style.name if style is not None else ""
            if style_name and re.fullmatch(r"Heading [1-9]", style_name):
                outline.append(DocumentLocation(
                    unit=len(units), kind="paragraph", text=block.text[:200], text_offset=0,
                    level=int(style_name[-1]) - 1,
                ))
        elif isinstance(block, Table):
            units.append(DocumentUnit(kind="table", text="\n".join(
                "\t".join(cell.text for cell in row.cells) for row in block.rows
            )))
        text_size += len(units[-1].text)
        if len(units) > 10_000 or text_size > 1_000_000:
            raise UserFacingError("文档超过处理上限，请拆分后再添加。", code="DOCUMENT_TOO_COMPLEX")
    return ParsedDocument(format, tuple(units), tuple(outline))


def _document(data: bytes) -> ParsedDocument:
    format = document_format(data)
    if format is None:
        raise UserFacingError("请选择 PDF、DOCX、XLSX 或 PPTX 文档。", code="DOCUMENT_UNSUPPORTED")
    try:
        return _parse(data, format, PARSER_VERSION)
    except (PdfReadError, BadZipFile, SyntaxError, DocxInvalidXmlError, PptxInvalidXmlError) as exc:
        raise UserFacingError("文档已损坏，无法读取。", code="DOCUMENT_INVALID") from exc


def _page(data: bytes, unit: int, offset: int, max_bytes: int, unit_count: int) -> DocumentPage:
    document = _document(data)
    units = document.units
    if unit > max(1, len(units)):
        raise UserFacingError("页码或段落编号超出文档范围。", code="DOCUMENT_RANGE_INVALID")
    end_unit = min(len(units), unit + unit_count - 1)
    selected = units[unit - 1:end_unit]
    kind = selected[0].kind if selected else ("page" if document.format == "pdf" else "paragraph")
    text = "\n\n".join(
        f"[{index} {item.kind}]\n{item.text}" if len(selected) > 1 else item.text
        for index, item in enumerate(selected, start=unit)
    )
    if offset > len(text):
        raise UserFacingError("页码或段落编号超出文档范围。", code="DOCUMENT_RANGE_INVALID")
    content = text[offset:].encode("utf-8")[:max_bytes].decode("utf-8", errors="ignore")
    end = offset + len(content)
    return DocumentPage(
        format=document.format, sha256=hashlib.sha256(data).hexdigest(), parser_version=PARSER_VERSION,
        unit=unit if units else 0, end_unit=end_unit, total_units=len(units), kind=kind,
        text=content, text_offset=offset, next_offset=end if end < len(text) else None,
        next_unit=end_unit + 1 if end_unit < len(units) else None,
    )


async def read_document(
    data: bytes, *, unit: int = 1, text_offset: int = 0, max_bytes: int = 16_000, unit_count: int = 1,
) -> DocumentPage:
    if len(data) > DOCUMENT_LIMIT:
        raise UserFacingError("文件不能超过 16 MiB。", code="ATTACHMENT_TOO_LARGE")
    if unit < 1 or text_offset < 0 or not 4 <= max_bytes <= 64_000 or not 1 <= unit_count <= 10:
        raise UserFacingError("页码或段落编号超出文档范围。", code="DOCUMENT_RANGE_INVALID")
    try:
        with fail_after(20):
            return await to_process.run_sync(_page, data, unit, text_offset, max_bytes, unit_count,
                                             cancellable=True, limiter=_parser_slots)
    except TimeoutError as exc:
        raise UserFacingError("文档处理超时，请拆分后重试。", code="DOCUMENT_TIMEOUT") from exc
    except BrokenWorkerProcess as exc:
        raise UserFacingError("文档处理失败，请检查文档或拆分后重试。", code="DOCUMENT_PARSE_FAILED") from exc


def _navigate(
    data: bytes, action: Literal["outline", "search"], query: str, offset: int, limit: int,
    ignore_case: bool,
) -> DocumentNavigation:
    document = _document(data)
    if action == "outline":
        items = list(document.outline[offset:offset + limit])
        next_offset = offset + limit if offset + limit < len(document.outline) else None
        return DocumentNavigation(
            format=document.format, sha256=hashlib.sha256(data).hexdigest(), parser_version=PARSER_VERSION,
            action=action, query="", total_units=len(document.units),
            empty_text_units=sum(not unit.text.strip() for unit in document.units),
            items=items, next_offset=next_offset, outline_truncated=document.outline_truncated,
        )
    # Escape user text: literal substring search cannot introduce arbitrary regex backtracking.
    # re.IGNORECASE preserves offsets into the original Unicode text, unlike casefolded copies.
    pattern = re.compile(re.escape(query), re.IGNORECASE if ignore_case else 0)
    matches: list[DocumentLocation] = []
    seen = 0
    for number, unit in enumerate(document.units, start=1):
        for match in pattern.finditer(unit.text):
            seen += 1
            if seen <= offset:
                continue
            if len(matches) == limit:
                break
            start = max(0, match.start() - 80)
            matches.append(DocumentLocation(
                unit=number, kind=unit.kind, text=unit.text[start:match.end() + 80], text_offset=start,
            ))
        if seen > offset + limit:
            break
    return DocumentNavigation(
        format=document.format, sha256=hashlib.sha256(data).hexdigest(), parser_version=PARSER_VERSION,
        action=action, query=query, total_units=len(document.units),
        empty_text_units=sum(not unit.text.strip() for unit in document.units),
        items=matches, next_offset=offset + limit if seen > offset + limit else None,
        outline_truncated=document.outline_truncated,
    )


async def navigate_document(
    data: bytes, *, action: Literal["outline", "search"], query: str = "", offset: int = 0,
    limit: int = 20, ignore_case: bool = True,
) -> DocumentNavigation:
    if len(data) > DOCUMENT_LIMIT:
        raise UserFacingError("文件不能超过 16 MiB。", code="ATTACHMENT_TOO_LARGE")
    if not 0 <= offset <= 1_000_000 or not 1 <= limit <= 20:
        raise UserFacingError("页码或段落编号超出文档范围。", code="DOCUMENT_RANGE_INVALID")
    if (action == "search" and not query.strip()) or len(query) > 200 or (action == "outline" and query):
        raise UserFacingError("请输入 1 至 200 个字符的搜索内容。", code="DOCUMENT_QUERY_INVALID")
    try:
        with fail_after(20):
            return await to_process.run_sync(
                _navigate, data, action, query, offset, limit, ignore_case,
                cancellable=True, limiter=_parser_slots,
            )
    except TimeoutError as exc:
        raise UserFacingError("文档处理超时，请拆分后重试。", code="DOCUMENT_TIMEOUT") from exc
    except BrokenWorkerProcess as exc:
        raise UserFacingError("文档处理失败，请检查文档或拆分后重试。", code="DOCUMENT_PARSE_FAILED") from exc
